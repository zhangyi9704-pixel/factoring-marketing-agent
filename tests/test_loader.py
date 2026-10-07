"""KnowledgeLoader 单元测试（不触网）。"""
import docx as docx_lib
import pytest

from app.rag.loader import KnowledgeLoader, LoaderError


class TestLoad:

    def test_md_utf8(self, tmp_path):
        f = tmp_path / "产品手册.md"
        f.write_text("# 保理产品\n正向保理说明。", encoding="utf-8")
        assert "正向保理" in KnowledgeLoader().load(f)

    def test_txt_gbk_fallback(self, tmp_path):
        f = tmp_path / "legacy.txt"
        f.write_text("商业保理GBK编码文档。", encoding="gbk")
        assert "商业保理" in KnowledgeLoader().load(f)

    def test_docx_paragraphs(self, tmp_path):
        document = docx_lib.Document()
        document.add_paragraph("段落一：产品要素。")
        document.add_paragraph("段落二：适用行业。")
        f = tmp_path / "手册.docx"
        document.save(str(f))
        text = KnowledgeLoader().load(f)
        assert "段落一" in text and "段落二" in text

    def test_missing_file(self, tmp_path):
        with pytest.raises(LoaderError, match="文件不存在"):
            KnowledgeLoader().load(tmp_path / "nope.md")

    def test_unsupported_suffix(self, tmp_path):
        f = tmp_path / "x.exe"
        f.write_bytes(b"MZ")
        with pytest.raises(LoaderError, match="不支持的文件类型"):
            KnowledgeLoader().load(f)

    def test_stripped_output(self, tmp_path):
        f = tmp_path / "a.txt"
        f.write_text("  \n 核心内容 \n ", encoding="utf-8")
        assert KnowledgeLoader().load(f) == "核心内容"


from app.rag.loader import parse_metadata
from test_ingest import HEADER


def test_metadata_normal_and_yaml_date(tmp_path):
    path = tmp_path / "正常.md"
    path.write_text(HEADER.format(version="KB-v1.0", status="conflicted") + "正文。", encoding="utf-8")
    body, meta = KnowledgeLoader().load_document(path)
    assert body == "正文。" and meta["effective_date"] == "2026-09-05"
    assert meta["status"] == "conflicted"


@pytest.mark.parametrize("field", ["doc_id", "category", "industry", "source", "version", "effective_date", "status", "topic"])
def test_metadata_missing_names_file(field):
    header = HEADER.format(version="v1", status="active")
    header = "\n".join(line for line in header.splitlines() if not line.startswith(field + ":"))
    with pytest.raises(LoaderError, match=f"坏.md.*{field}"):
        parse_metadata(header + "\n正文", "坏.md")


@pytest.mark.parametrize("old,new", [("category: product", "category: wrong"), ("industry: 基建", "industry: other"), ("status: active", "status: indexed"), ("2026-09-05", "2026-02-30")])
def test_metadata_invalid_enum_or_date(old, new):
    with pytest.raises(LoaderError, match="错误.md"):
        parse_metadata(HEADER.format(version="v1", status="active").replace(old, new) + "正文", "错误.md")
