"""RAG 文档加载（TDD §11.1 第一步：pdf/docx/md/txt → 纯文本）。

- md / txt：utf-8 优先，解码失败回退 gbk，再失败以替换字符兜底；
- pdf：pypdf 逐页抽取；
- docx：python-docx 段落 + 表格（表格行以 " | " 连接）。

切分职责在 chunker（parent_child），本模块只负责"文件 → 文本"。
"""
from __future__ import annotations

from pathlib import Path
from typing import Union
from datetime import date
import yaml

METADATA_FIELDS = ("doc_id", "category", "industry", "source", "version",
                   "effective_date", "status", "topic")


def parse_metadata(text: str, filename: str) -> tuple[str, dict[str, str]]:
    """严格解析条目头；正文和知识治理字段分离，错误包含文件名。"""
    lines = text.lstrip("\ufeff").splitlines()
    if not lines or lines[0].strip() != "---":
        raise LoaderError(f"{filename}: 缺少 YAML 元数据头")
    try:
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        meta = yaml.safe_load("\n".join(lines[1:end]))
    except (StopIteration, yaml.YAMLError, ValueError) as exc:
        raise LoaderError(f"{filename}: YAML 元数据格式错误") from exc
    if not isinstance(meta, dict):
        raise LoaderError(f"{filename}: 元数据必须是对象")
    missing = [key for key in METADATA_FIELDS if key not in meta or meta[key] is None
               or not str(meta[key]).strip()]
    if missing:
        raise LoaderError(f"{filename}: 缺少必填字段 {', '.join(missing)}")
    for key in METADATA_FIELDS:
        if not isinstance(meta[key], str) and not (key == "effective_date" and isinstance(meta[key], date)):
            raise LoaderError(f"{filename}: {key} 必须是字符串")
    meta = {key: str(meta[key]).strip() for key in METADATA_FIELDS}
    for key, allowed in {"category": {"product", "process", "industry", "compliance"},
                         "industry": {"all", "基建", "医药", "电商"},
                         "status": {"active", "conflicted", "archived"}}.items():
        if meta[key] not in allowed:
            raise LoaderError(f"{filename}: {key} 枚举错误")
    try:
        if date.fromisoformat(meta["effective_date"]).isoformat() != meta["effective_date"]:
            raise ValueError()
    except ValueError as exc:
        raise LoaderError(f"{filename}: effective_date 必须是 YYYY-MM-DD") from exc
    body = "\n".join(lines[end + 1:]).strip()
    if not body:
        raise LoaderError(f"{filename}: 正文为空")
    return body, meta


class LoaderError(RuntimeError):
    """文档解析错误。"""


class KnowledgeLoader:
    """知识库文档加载器。"""

    SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}

    def load_document(self, path: Union[str, Path]) -> tuple[str, dict[str, str]]:
        return parse_metadata(self.load(path), Path(path).name)

    def load(self, path: Union[str, Path]) -> str:
        """读取单个文档，返回纯文本（首尾空白已去除）。"""
        path = Path(path)
        if not path.is_file():
            raise LoaderError(f"文件不存在: {path}")
        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED_SUFFIXES:
            raise LoaderError(
                f"不支持的文件类型 {suffix!r}"
                f"（支持 {'/'.join(sorted(self.SUPPORTED_SUFFIXES))}）: {path}"
            )
        try:
            if suffix in {".md", ".txt"}:
                return self._load_text(path)
            if suffix == ".pdf":
                return self._load_pdf(path)
            return self._load_docx(path)
        except LoaderError:
            raise
        except Exception as exc:
            raise LoaderError(f"解析失败 {path}: {exc}") from exc

    # -------------------------------------------------
    # 各格式解析
    # -------------------------------------------------

    @staticmethod
    def _decode(raw: bytes) -> str:
        for encoding in ("utf-8", "gbk"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    def _load_text(self, path: Path) -> str:
        return self._decode(path.read_bytes()).strip()

    @staticmethod
    def _load_pdf(path: Path) -> str:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = (page.extract_text() or "" for page in reader.pages)
        return "\n\n".join(p.strip() for p in pages if p and p.strip()).strip()

    @staticmethod
    def _load_docx(path: Path) -> str:
        import docx

        document = docx.Document(str(path))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts).strip()
